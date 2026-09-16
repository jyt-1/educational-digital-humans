# [工单17] 人工智能NLP-Agent数字人项目-教育智能体-智能备课任务 —— 文档导出服务
"""导出服务：python-docx 导出教案/习题/试题/案例；python-pptx 导出课件。

验收要求：导出内容必须与前端编辑后的内容一致（设计文档 2.2 场景一）。
"""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches, Pt as PptxPt

# 导出文档页脚统一标识（设计文档 5.4 节：AI 生成内容须标注）
AI_NOTICE = "本文件由 AI 辅助生成，已由教师复核后导出。"


# ============================================================ Markdown -> docx

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")


def _add_runs_with_bold(paragraph, text: str) -> None:
    """把 **加粗** 语法转成 docx 的粗体 run。"""
    pos = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _parse_table_row(line: str) -> list[str]:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def _flush_table(doc: Document, rows: list[list[str]]) -> None:
    """把收集到的 Markdown 表格写入 docx。"""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell_text = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            cell.text = ""
            para = cell.paragraphs[0]
            _add_runs_with_bold(para, cell_text)
            if i == 0:
                for run in para.runs:
                    run.bold = True
    doc.add_paragraph()


def markdown_to_docx(markdown_text: str, title: str, subtitle: str | None = None) -> bytes:
    """把 Markdown 文本渲染为 docx 字节流。支持标题、列表、加粗、表格、引用。"""
    doc = Document()

    doc.add_heading(title, level=0)
    if subtitle:
        para = doc.add_paragraph(subtitle)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.runs[0].font.size = Pt(10)

    pending_table: list[list[str]] = []

    for raw_line in (markdown_text or "").splitlines():
        line = raw_line.rstrip()

        # 表格行收集
        if line.strip().startswith("|"):
            if _TABLE_SEP_RE.match(line):
                continue  # 分隔行 |---|---|
            pending_table.append(_parse_table_row(line))
            continue
        if pending_table:
            _flush_table(doc, pending_table)
            pending_table = []

        if not line.strip():
            continue

        # 标题
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = min(len(heading.group(1)), 4)
            doc.add_heading(heading.group(2).strip(), level=level)
            continue

        # 无序列表
        bullet = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if bullet:
            para = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(para, bullet.group(1))
            continue

        # 有序列表
        ordered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if ordered:
            para = doc.add_paragraph(style="List Number")
            _add_runs_with_bold(para, ordered.group(1))
            continue

        # 引用
        quote = re.match(r"^\s*>\s?(.*)$", line)
        if quote:
            para = doc.add_paragraph(style="Intense Quote")
            _add_runs_with_bold(para, quote.group(1))
            continue

        # 普通段落
        para = doc.add_paragraph()
        _add_runs_with_bold(para, line)

    if pending_table:
        _flush_table(doc, pending_table)

    # 页脚标识（AI 生成内容标注义务）
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(AI_NOTICE)
    run.font.size = Pt(8)
    run.italic = True

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ============================================================ 习题/试题 -> docx

def questions_to_docx(
    items: list[dict],
    title: str,
    subtitle: str | None = None,
    with_answer: bool = True,
) -> bytes:
    """把习题/试题列表渲染为 docx。"""
    doc = Document()
    doc.add_heading(title, level=0)
    if subtitle:
        para = doc.add_paragraph(subtitle)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for idx, item in enumerate(items, start=1):
        qtype = item.get("qtype") or ""
        score = item.get("score")
        head = f"{idx}. [{qtype}]" + (f"（{score} 分）" if score else "")
        para = doc.add_paragraph()
        run = para.add_run(head)
        run.bold = True

        stem = str(item.get("stem") or "").strip()
        if stem:
            doc.add_paragraph(stem)

        options = item.get("options")
        if isinstance(options, list):
            for opt in options:
                doc.add_paragraph(str(opt), style="List Bullet")

        if with_answer:
            answer = item.get("answer")
            if answer:
                p = doc.add_paragraph()
                p.add_run("【答案】").bold = True
                p.add_run(str(answer))
            analysis = item.get("analysis")
            if analysis:
                p = doc.add_paragraph()
                p.add_run("【解析】").bold = True
                p.add_run(str(analysis))
            kp = item.get("knowledge_point")
            if kp:
                p = doc.add_paragraph()
                p.add_run("【知识点】").bold = True
                p.add_run(str(kp))

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(AI_NOTICE)
    run.font.size = Pt(8)
    run.italic = True

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ============================================================ 课件 -> pptx

def courseware_to_pptx(slides: list[dict], title: str) -> bytes:
    """把课件 JSON（[{title, bullets, notes}]）渲染为 pptx。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 首页封面
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = title
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = "AI 辅助生成 · 教师复核后使用"

    for slide_data in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide_title = str(slide_data.get("title") or "").strip() or "未命名"
        slide.shapes.title.text = slide_title

        bullets = slide_data.get("bullets") or []
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(bullet)
            para.font.size = PptxPt(20)

        # 备注写入演讲者备注
        notes = slide_data.get("notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = str(notes)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


# ============================================================ 统一入口

def export_plan(plan_content: dict[str, Any], content_type: str, title: str, subtitle: str | None) -> tuple[bytes, str, str]:
    """按内容类型导出，返回 (文件字节, 文件名, media_type)。"""
    items = plan_content.get("items") or []
    raw = plan_content.get("raw") or ""

    if content_type == "课件":
        data = courseware_to_pptx(items, title)
        return (
            data,
            f"{title}.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

    if content_type in ("习题", "试题"):
        data = questions_to_docx(items, title, subtitle)
        return (
            data,
            f"{title}.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    # 教案 / 案例：Markdown 正文
    data = markdown_to_docx(raw, title, subtitle)
    return (
        data,
        f"{title}.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
