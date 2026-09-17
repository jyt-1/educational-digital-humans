# [工单18] 人工智能NLP-Agent数字人项目-教育智能体-智能助教任务 —— XLS/XLSX 解析器
"""表格文档解析（openpyxl）：每个工作表转成一张 Markdown 表格。

多模态处理（工单18）：表格统一转 Markdown 存储，引用时按「工作表名 + 页码」溯源。
为避免超大表格撑爆向量库，单表限制行/列上限，超出部分截断并在表前注明。
"""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import load_workbook

from app.services.parsers.base import BlockBuilder, ParsedBlock, table_to_markdown

logger = logging.getLogger(__name__)

_MAX_ROWS = 300  # 单表最多保留行数
_MAX_COLS = 30  # 单表最多保留列数


def parse(path: Path, image_dir: Path | None = None, rel_prefix: str = "") -> list[ParsedBlock]:
    """解析 Excel，每个工作表输出一个 table 块（page_no = 工作表序号）。"""
    builder = BlockBuilder()
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        for sheet_no, sheet in enumerate(workbook.worksheets, start=1):
            rows, truncated_rows = _read_rows(sheet)
            if not rows:
                continue
            truncated_cols = max(len(r) for r in rows) >= _MAX_COLS
            markdown = table_to_markdown(rows)
            if not markdown:
                continue

            header = f"### 工作表：{sheet.title}"
            if truncated_rows or truncated_cols:
                header += (
                    f"（仅保留前 {_MAX_ROWS} 行 × {_MAX_COLS} 列，"
                    "完整数据请下载原文件查看）"
                )
            builder.add(
                ParsedBlock(
                    "table",
                    content=f"{header}\n\n{markdown}",
                    page_no=sheet_no,
                    section=sheet.title,
                )
            )
    finally:
        workbook.close()

    return builder.blocks


def _read_rows(sheet) -> tuple[list[list[str]], bool]:
    rows: list[list[str]] = []
    truncated = False
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        if index >= _MAX_ROWS:
            truncated = True
            break
        cells = ["" if value is None else str(value) for value in row[:_MAX_COLS]]
        if any(cell.strip() for cell in cells):
            rows.append(cells)
    return rows, truncated
