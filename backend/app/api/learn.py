# [工单19] 人工智能NLP-Agent数字人项目-教育智能体-个性化学习推荐任务 —— 个性化学习接口
"""工单19 接口层：知识图谱、学生画像、学习路径、自适应练习、试卷模式、AIGC 错题本、联动助教。

接口清单与权限见设计文档 3.2.2「（3）个性化学习推荐」。三条贯穿全模块的口径：

1. **「未归类」占位节点在画像、路径、练习中一律过滤**（设计文档 2.2 场景三第 8 条）。
   它的掌握度没有解释力，出现在雷达图上就是一个永远不满、也无法补救的轴。
2. **题库为空必须给明确空态，不得静默返回空列表**。空态分三种（没同步 / 没学习记录 /
   该知识点没题），三者的处置完全不同——合成一条"暂无数据"，用户不知道该做什么。
3. **写接口按 `user_id` 隔离**：画像、错题本、成绩导入都是个人数据。
"""

from __future__ import annotations

import io
import logging
import uuid
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.auth import get_current_user, require_role
from app.config import settings
from app.db import get_db
from app.models.assistant import Conversation, Message
from app.models.learn import (
    ANALYZE_DONE,
    SOURCE_EXAM,
    SOURCE_MANUAL,
    SOURCE_SCORE_IMPORT,
    Attempt,
    KnowledgePoint,
    KpAlias,
    KpFaq,
    MessageKp,
    MistakeBook,
    Question,
    ScoreImport,
)
from app.models.lesson import ExamQuestion, TeachingPlan
from app.models.user import User
from app.schemas.common import ApiResponse
from app.schemas.learn import (
    AnswerRequest,
    ExamSubmitRequest,
    KpMergeRequest,
    ScoreImportRequest,
    VariantAnswerRequest,
)
from app.services import learn_profile, learn_sync, mistake as mistake_service, retriever
from app.services.kp_match import ON_UNMATCHED_DROP, load_index, match_label, normalize_label
from app.services.learn_answer import ensure_mistake, load_options, record_attempt

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/learn", tags=["个性化学习推荐"])

require_teacher = require_role("teacher")
require_student = require_role("student")

# 空态原因。前端据 reason 给出不同引导，而不是一律显示"暂无数据"
REASON_BANK_EMPTY = "bank_empty"
REASON_NO_TARGET = "no_target"
REASON_NO_QUESTION_FOR_KP = "no_question_for_kp"

# 单次取题上限。防止前端传个 count=10000 把整个题库捞出来
MAX_PRACTICE_COUNT = 20
RELATED_QUESTION_LIMIT = 5
RELATED_MATERIAL_LIMIT = 3


# ------------------------------------------------------------------ 工具

def _question_brief(question: Question, *, with_answer: bool = False) -> dict:
    """题目的对外结构。`with_answer=False` 时**不带答案与解析**——
    练习与试卷都要学生先作答，答案由提交接口返回。"""
    payload = {
        "question_id": question.id,
        "kp_id": question.kp_id,
        "qtype": question.qtype,
        "stem": question.stem,
        "options": load_options(question.options_json),
        "difficulty": question.difficulty,
    }
    if with_answer:
        payload["answer"] = question.answer
        payload["analysis"] = question.analysis
    return payload


def _kp_names(db: Session) -> dict[int, str]:
    return {node.id: node.name for node in db.execute(select(KnowledgePoint)).scalars()}


def _bank_is_empty(db: Session) -> bool:
    return not db.scalar(select(func.count()).select_from(Question))


# ------------------------------------------------------------------ 知识图谱

@router.get("/knowledge-points", summary="知识点图谱（含先修关系）")
def list_knowledge_points(
    course: str | None = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """图谱展示用。**过滤 `is_placeholder`**：占位节点不是一个知识点。"""
    query = (
        select(KnowledgePoint)
        .where(KnowledgePoint.is_placeholder == 0)
        .order_by(KnowledgePoint.order_no)
    )
    if course:
        query = query.where(KnowledgePoint.course == course)
    nodes = list(db.execute(query).scalars())
    return ApiResponse.ok(
        {
            "total": len(nodes),
            "items": [
                {
                    "kp_id": node.id,
                    "name": node.name,
                    "course": node.course,
                    "prereq_id": node.prereq_id,
                    "order_no": node.order_no,
                }
                for node in nodes
            ],
        }
    )


# ------------------------------------------------------------------ 画像与路径

@router.get("/profile", summary="学生画像（掌握度雷达图数据）")
def get_profile(
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """**实时计算优先、缓存兜底**（设计文档 3.3.3）：缓存缺失或过期时现算并回写，
    任何情况下都不返回空画像。"""
    profile = learn_profile.read_profile(db, user.id)
    nodes = {
        node.id: node
        for node in db.execute(
            select(KnowledgePoint).where(KnowledgePoint.is_placeholder == 0)
        ).scalars()
    }

    items = [
        {
            "kp_id": kp_id,
            "name": nodes[kp_id].name,
            "order_no": nodes[kp_id].order_no,
            "prereq_id": nodes[kp_id].prereq_id,
            **record.to_dict(),
        }
        for kp_id, record in profile.items()
        if kp_id in nodes
    ]
    items.sort(key=lambda item: item["order_no"])

    mastery_values = [item["mastery"] for item in items]
    weak = [item for item in items if item["mastery"] < learn_profile.WEAK_THRESHOLD]
    return ApiResponse.ok(
        {
            "items": items,
            "kp_count": len(items),
            "average_mastery": round(sum(mastery_values) / len(mastery_values), 4)
            if mastery_values
            else 0.0,
            "weak_count": len(weak),
            "weak_threshold": learn_profile.WEAK_THRESHOLD,
        }
    )


@router.get("/path", summary="学习路径推荐")
def get_path(
    course: str | None = Query(None),
    limit: int | None = Query(None, ge=1, le=50, description="仪表盘取前 N 条摘要"),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """每个节点由前端挂 `/related` 三个出口（设计文档 3.2.2 v1.3 补充②）。"""
    items = learn_profile.recommend_path(db, user.id, course=course, limit=limit)
    return ApiResponse.ok(
        {
            "items": items,
            "total": len(items),
            "weak_threshold": learn_profile.WEAK_THRESHOLD,
            "message": None
            if items
            else "暂无薄弱知识点。完成一次练习或导入历史成绩后，这里会给出推荐学习顺序。",
        }
    )


# ------------------------------------------------------------------ 自适应练习

@router.get("/practice", summary="自适应练习取题")
def get_practice(
    kp_id: int | None = Query(None, description="缺省取推荐路径的第一个知识点"),
    count: int = Query(5, ge=1, le=MAX_PRACTICE_COUNT),
    db: Session = Depends(get_db),
    user: User = Depends(require_student),
):
    if _bank_is_empty(db):
        return ApiResponse.ok(
            {
                "questions": [],
                "reason": REASON_BANK_EMPTY,
                "message": "题库尚未同步，请先执行 POST /api/learn/kp/sync?target=questions（需教师身份）",
            }
        )

    node = None
    if kp_id is not None:
        node = db.scalar(
            select(KnowledgePoint).where(
                KnowledgePoint.id == kp_id, KnowledgePoint.is_placeholder == 0
            )
        )
        if node is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="知识点不存在")
    else:
        path = learn_profile.recommend_path(db, user.id, limit=1)
        if not path:
            return ApiResponse.ok(
                {
                    "questions": [],
                    "reason": REASON_NO_TARGET,
                    "message": "还没有可推荐的知识点。先做一次练习或导入历史成绩，系统才能定位你的薄弱点。",
                }
            )
        node = db.get(KnowledgePoint, path[0]["kp_id"])

    state = learn_profile.get_practice_state(db, user.id, node.id)
    db.commit()

    questions = learn_profile.pick_questions(
        db, kp_id=node.id, difficulty=state.difficulty, count=count
    )
    if not questions:
        return ApiResponse.ok(
            {
                "questions": [],
                "kp_id": node.id,
                "kp_name": node.name,
                "difficulty": state.difficulty,
                "reason": REASON_NO_QUESTION_FOR_KP,
                "message": f"知识点「{node.name}」下暂无题目，请让教师重新同步题库。",
            }
        )

    return ApiResponse.ok(
        {
            "kp_id": node.id,
            "kp_name": node.name,
            "difficulty": state.difficulty,
            "streak_correct": state.streak_correct,
            "promote_streak": learn_profile.PROMOTE_STREAK,
            "questions": [_question_brief(question) for question in questions],
            "reason": None,
            "message": None,
        }
    )


@router.post("/answer", summary="提交练习作答")
def submit_answer(
    body: AnswerRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_student),
):
    """判分 → 写 `attempts` → 推进难度档 → 答错进错题本 → 刷新画像缓存。

    `practice_state` 只由本接口推进（见 `learn_answer` 模块注释）。
    """
    question = db.get(Question, body.question_id)
    if question is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="题目不存在")

    attempt, is_correct = record_attempt(
        db, student_id=user.id, question=question, user_answer=body.user_answer
    )

    difficulty = None
    streak = 0
    mastery = None
    if question.kp_id is not None:
        state = learn_profile.get_practice_state(db, user.id, question.kp_id)
        difficulty = learn_profile.apply_answer_to_state(state, is_correct)
        streak = state.streak_correct

    mistake_id = None
    if not is_correct:
        mistake = ensure_mistake(db, student_id=user.id, question=question, attempt=attempt)
        mistake_id = mistake.id

    profile = learn_profile.refresh_profile_cache(db, user.id)
    if question.kp_id is not None and question.kp_id in profile:
        mastery = round(profile[question.kp_id].mastery, 4)

    return ApiResponse.ok(
        {
            "attempt_id": attempt.id,
            "is_correct": is_correct,
            "correct_answer": question.answer,
            "analysis": question.analysis,
            "difficulty": difficulty,
            "streak_correct": streak,
            "promote_streak": learn_profile.PROMOTE_STREAK,
            "mastery": mastery,
            "mistake_id": mistake_id,
        }
    )


# ------------------------------------------------------------------ 试卷模式

def _resolve_exam_plan(db: Session, plan_id: int | None) -> TeachingPlan | None:
    """定位一套试题：给定 `plan_id` 就用它，否则取**最近一次含试题的备课记录**。"""
    if plan_id is not None:
        return db.get(TeachingPlan, plan_id)
    plans_with_exam = (
        select(ExamQuestion.plan_id).where(ExamQuestion.plan_id.isnot(None)).distinct()
    )
    return db.scalar(
        select(TeachingPlan)
        .where(TeachingPlan.id.in_(plans_with_exam))
        .order_by(TeachingPlan.created_at.desc(), TeachingPlan.id.desc())
    )


def _exam_rows(db: Session, plan_id: int) -> list[tuple[Question, ExamQuestion]]:
    """同源整套试题。分值存在 `exam_questions.score` 上——`questions` 表按设计没有分值列
    （见 3.3.2 建表 SQL），故必须回连源表取，不能想当然地去 `questions` 里找。"""
    return list(
        db.execute(
            select(Question, ExamQuestion)
            .join(ExamQuestion, ExamQuestion.id == Question.source_id)
            .where(Question.source == SOURCE_EXAM, ExamQuestion.plan_id == plan_id)
            .order_by(Question.id)
        ).all()
    )


@router.get("/exam", summary="试卷模式·组卷")
def get_exam(
    plan_id: int | None = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    plan = _resolve_exam_plan(db, plan_id)
    if plan is None:
        return ApiResponse.ok(
            {
                "questions": [],
                "reason": REASON_BANK_EMPTY,
                "message": "还没有可用的试卷。请先在「智能备课」中生成一套试题，并执行一次题库同步。",
            }
        )

    rows = _exam_rows(db, plan.id)
    if not rows:
        return ApiResponse.ok(
            {
                "plan_id": plan.id,
                "questions": [],
                "reason": REASON_BANK_EMPTY,
                "message": "该套试题尚未同步到题库，请先执行 POST /api/learn/kp/sync?target=questions（需教师身份）",
            }
        )

    return ApiResponse.ok(
        {
            "plan_id": plan.id,
            "title": plan.title or f"{plan.course_name} 试题",
            "course_name": plan.course_name,
            "chapter": plan.chapter,
            "question_count": len(rows),
            "total_score": sum(exam.score or 0 for _question, exam in rows),
            "questions": [
                {**_question_brief(question), "score": exam.score or 0}
                for question, exam in rows
            ],
            "reason": None,
            "message": None,
        }
    )


@router.post("/exam/submit", summary="试卷模式·交卷")
def submit_exam(
    body: ExamSubmitRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_student),
):
    """批量判分 → 逐题写 `attempts`（`is_exam=1`）→ 刷新画像。**不动 `practice_state`**。

    **未作答的题也写一条 `is_correct=0` 的作答**：跳过就是没拿到分，画像理应反映这一点；
    若不写，一场"交白卷"对画像毫无影响，这与"考试计入画像"的设计不符。
    """
    plan = _resolve_exam_plan(db, body.plan_id)
    if plan is None or plan.id != body.plan_id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="试卷不存在")

    rows = _exam_rows(db, plan.id)
    if not rows:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="该试卷尚未同步到题库")

    answers = {item.question_id: item.user_answer for item in body.answers}
    allowed = {question.id for question, _exam in rows}
    unknown = set(answers) - allowed
    if unknown:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"提交了不属于本试卷的题目：{sorted(unknown)}",
        )

    items = []
    earned_total = 0
    full_total = 0
    correct_count = 0
    for question, exam in rows:
        user_answer = answers.get(question.id)
        _attempt, is_correct = record_attempt(
            db, student_id=user.id, question=question, user_answer=user_answer, is_exam=True
        )
        score = exam.score or 0
        earned = score if is_correct else 0
        earned_total += earned
        full_total += score
        correct_count += 1 if is_correct else 0
        items.append(
            {
                "question_id": question.id,
                "is_correct": is_correct,
                "user_answer": user_answer,
                "correct_answer": question.answer,
                "analysis": question.analysis,
                "score": score,
                "earned": earned,
            }
        )

    learn_profile.refresh_profile_cache(db, user.id)

    return ApiResponse.ok(
        {
            "plan_id": plan.id,
            "title": plan.title or f"{plan.course_name} 试题",
            "score": earned_total,
            "total_score": full_total,
            "correct_count": correct_count,
            "question_count": len(rows),
            "items": items,
        }
    )


# ------------------------------------------------------------------ AIGC 错题本

def _get_own_mistake(db: Session, mistake_id: int, user: User) -> MistakeBook:
    """取本人错题。越权与不存在统一 404，不泄露他人数据是否存在。"""
    row = db.scalar(
        select(MistakeBook).where(
            MistakeBook.id == mistake_id, MistakeBook.student_id == user.id
        )
    )
    if row is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="错题不存在")
    return row


def _mistake_detail(db: Session, mistake: MistakeBook) -> dict:
    question = db.get(Question, mistake.question_id)
    kp = db.get(KnowledgePoint, mistake.kp_id) if mistake.kp_id else None
    attempt = db.scalar(
        select(Attempt)
        .where(Attempt.id == mistake.attempt_id)
        .order_by(Attempt.id.desc())
    )
    if attempt is None:
        attempt = db.scalar(
            select(Attempt)
            .where(
                Attempt.student_id == mistake.student_id,
                Attempt.question_id == mistake.question_id,
                Attempt.is_correct.is_(False),
            )
            .order_by(Attempt.id.desc())
        )
    return {
        "mistake_id": mistake.id,
        "question_id": mistake.question_id,
        "kp_id": mistake.kp_id,
        "kp_name": kp.name if kp else None,
        "qtype": question.qtype if question else None,
        "stem": question.stem if question else None,
        "options": load_options(question.options_json) if question else [],
        "user_answer": attempt.user_answer if attempt else None,
        "correct_answer": question.answer if question else None,
        "analysis_status": mistake.analyze_status,
        "analysis": mistake_service.load_analysis(mistake),
        "variant_questions": mistake_service.load_variants(mistake),
        "variant_tries": mistake.variant_tries,
        "created_at": mistake.created_at.isoformat() if mistake.created_at else None,
    }


@router.get("/mistakes", summary="错题本列表")
def list_mistakes(
    kp_id: int | None = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """只返回题干、知识点、状态与时间；**完整分析走详情接口**——
    单条分析含 2~3 道带解析的变式题，放进列表会让响应体随错题数线性膨胀。"""
    query = (
        select(MistakeBook, Question, KnowledgePoint.name)
        .join(Question, Question.id == MistakeBook.question_id)
        .outerjoin(KnowledgePoint, KnowledgePoint.id == MistakeBook.kp_id)
        .where(MistakeBook.student_id == user.id)
        .order_by(MistakeBook.created_at.desc(), MistakeBook.id.desc())
    )
    if kp_id is not None:
        query = query.where(MistakeBook.kp_id == kp_id)

    rows = db.execute(query).all()
    return ApiResponse.ok(
        {
            "total": len(rows),
            "items": [
                {
                    "mistake_id": mistake.id,
                    "question_id": mistake.question_id,
                    "kp_id": mistake.kp_id,
                    "kp_name": kp_name,
                    "stem": question.stem,
                    "difficulty": question.difficulty,
                    "analysis_status": mistake.analyze_status,
                    "variant_tries": mistake.variant_tries,
                    "created_at": mistake.created_at.isoformat() if mistake.created_at else None,
                }
                for mistake, question, kp_name in rows
            ],
        }
    )


@router.get("/mistakes/{mistake_id}", summary="错题详情")
def get_mistake(
    mistake_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    return ApiResponse.ok(_mistake_detail(db, _get_own_mistake(db, mistake_id, user)))


@router.post("/mistakes/{mistake_id}/analyze", summary="触发/重新触发 AIGC 错题分析")
async def analyze_mistake(
    mistake_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_student),
):
    """同步调用 LLM 并等待结果（本工单不做异步任务队列）。

    失败时**不把错误吞掉**——把 `analyze_status` 置 `failed` 并返回可读原因，
    否则前端只会显示一个空解析，现场无从解释（设计文档 3.2.7 只要求"有输出"，
    但"有输出"的前提是失败了看得见）。
    """
    mistake = _get_own_mistake(db, mistake_id, user)
    try:
        result = await mistake_service.analyze_mistake(db, mistake)
    except Exception as exc:  # noqa: BLE001 - LLM 侧任何异常都要变成可读提示
        logger.warning("错题 %s 分析失败：%s", mistake_id, exc)
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"错题分析失败：{exc}",
        ) from exc
    return ApiResponse.ok({"mistake_id": mistake.id, **result, "analysis_status": ANALYZE_DONE})


@router.post("/mistakes/{mistake_id}/variant-answer", summary="变式题作答")
async def answer_variant(
    mistake_id: int,
    body: VariantAnswerRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_student),
):
    """变式题判分；**答错则重新分析**（工单原文"变式题可再答，再错则重新分析"）。

    变式题以 `questions` 行的形式落库（见 `services/mistake.py` 模块注释），
    因此这里能写出带 `is_variant=1` 的 `attempts`——画像也就把变式练习一并计入。
    """
    mistake = _get_own_mistake(db, mistake_id, user)
    variants = mistake_service.load_variants(mistake)
    variant_ids = {item.get("question_id") for item in variants}
    if body.question_id not in variant_ids:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="该题不属于此错题的变式题"
        )

    question = db.get(Question, body.question_id)
    if question is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="题目不存在")

    _attempt, is_correct = record_attempt(
        db, student_id=user.id, question=question, user_answer=body.user_answer, is_variant=True
    )

    reanalyzed = False
    error = None
    if not is_correct:
        mistake.variant_tries += 1
        db.commit()
        try:
            await mistake_service.analyze_mistake(db, mistake)
            reanalyzed = True
        except Exception as exc:  # noqa: BLE001 - 分析失败不该让"作答已记录"这个事实丢失
            error = str(exc)
            logger.warning("错题 %s 变式题再错后重新分析失败：%s", mistake_id, exc)

    learn_profile.refresh_profile_cache(db, user.id)

    return ApiResponse.ok(
        {
            "is_correct": is_correct,
            "correct_answer": question.answer,
            "analysis": question.analysis,
            "variant_tries": mistake.variant_tries,
            "reanalyzed": reanalyzed,
            "error": error,
            "mistake": _mistake_detail(db, mistake),
        }
    )


# ------------------------------------------------------------------ 联动助教

@router.get("/related", summary="按知识点推荐（相关提问 / 学习资料 / 练习题）")
async def get_related(
    kp_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """三个入口共用：错题详情页侧栏、仪表盘今日任务、学习路径页每个节点。

    进入前先对**该生**做一次惰性补挂：只扫该生尚无 `message_kp` 的消息并即时挂靠，
    量极小。不这样做的话，"刚在助教里问过 → 当场看侧栏为空"必须等教师跑一次批处理
    `/kp/sync?target=messages` 才会出现，演示时看起来就像坏了一样。
    """
    kp = db.scalar(
        select(KnowledgePoint).where(
            KnowledgePoint.id == kp_id, KnowledgePoint.is_placeholder == 0
        )
    )
    if kp is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="知识点不存在")

    learn_sync.sync_messages(db, student_id=user.id)

    # ① 相关提问：两路都按该 kp_id 取，而不是取该生的全部提问
    own_rows = db.execute(
        select(Message.content, Message.conversation_id)
        .join(MessageKp, MessageKp.message_id == Message.id)
        .join(Conversation, Conversation.id == Message.conversation_id)
        .where(
            MessageKp.kp_id == kp_id,
            Message.role == "user",
            Conversation.user_id == user.id,
        )
        .order_by(Message.id.desc())
        .limit(RELATED_QUESTION_LIMIT)
    ).all()
    questions = [
        {"question": content, "source": "own", "conversation_id": conversation_id}
        for content, conversation_id in own_rows
    ]

    faq_query = select(KpFaq).where(KpFaq.kp_id == kp_id)
    if not settings.ASSISTANT_HOTQ_ENABLED:
        # 开关关闭时高频栏退化为只显示种子问题（设计文档 2.2 场景三第 7 条）
        faq_query = faq_query.where(KpFaq.source == "seed")
    faq_rows = list(
        db.execute(
            faq_query.order_by(KpFaq.ask_count.desc(), KpFaq.id).limit(RELATED_QUESTION_LIMIT)
        ).scalars()
    )
    questions.extend(
        # 跨用户聚合来的问题**没有会话可跳，也不能暴露提问者**：前端据 source
        # 决定是"跳回自己的原会话"还是"带问题文本去问答页预填发起新提问"
        {"question": row.question, "source": row.source, "conversation_id": None}
        for row in faq_rows
    )

    # ② 关联学习资料：走工单18 的公开检索入口，由它内部完成库隔离与整条检索链
    materials: list[dict] = []
    try:
        hits = await retriever.search(
            db, question=kp.name, user=user, top_k=RELATED_MATERIAL_LIMIT, scope="all"
        )
        materials = [hit.to_citation(index) for index, hit in enumerate(hits, start=1)]
    except Exception as exc:  # noqa: BLE001 - 检索侧异常不该让整条推荐失败
        logger.warning("知识点 %s 的资料检索失败：%s", kp.name, exc)

    # ③ 同知识点练习题：不带答案（走练习接口作答）
    practice_rows = list(
        db.execute(
            select(Question).where(Question.kp_id == kp_id).order_by(Question.id).limit(5)
        ).scalars()
    )

    return ApiResponse.ok(
        {
            "kp_id": kp.id,
            "kp_name": kp.name,
            "questions": questions,
            "materials": materials,
            "exercises": [
                {
                    "question_id": row.id,
                    "stem": row.stem,
                    "qtype": row.qtype,
                    "difficulty": row.difficulty,
                }
                for row in practice_rows
            ],
            "hotq_enabled": settings.ASSISTANT_HOTQ_ENABLED,
        }
    )


# ------------------------------------------------------------------ 历史成绩导入

TEMPLATE_HEADERS = ("知识点", "得分", "满分")


def _import_rows(
    db: Session,
    *,
    user: User,
    rows: list[tuple[str, float, float]],
    source_file: str | None,
    batch_id: str | None,
) -> dict:
    """导入管线：逐行校验 → 知识点匹配 → 写 `score_imports`。

    **匹配失败的行不入库**，列进失败明细让用户改文件重传——混进「未归类」会直接
    脏掉初始画像，而画像正是整个推荐链路的起点（设计文档 2.2 场景三第 1 条）。
    """
    index = load_index(db)
    batch = batch_id or str(uuid.uuid4())
    success: list[dict] = []
    failures: list[dict] = []

    for row_no, (raw_kp, score, total) in enumerate(rows, start=2):  # 第 1 行是表头
        label = (raw_kp or "").strip()
        if not label:
            failures.append({"row": row_no, "raw_label": "", "reason": "知识点为空"})
            continue
        if total is None or total <= 0:
            failures.append({"row": row_no, "raw_label": label, "reason": "满分必须大于 0"})
            continue
        if score is None or score < 0:
            failures.append({"row": row_no, "raw_label": label, "reason": "得分不能为负数"})
            continue
        if score > total:
            failures.append(
                {"row": row_no, "raw_label": label, "reason": f"得分 {score} 大于满分 {total}"}
            )
            continue

        result = match_label(label, index=index, on_unmatched=ON_UNMATCHED_DROP)
        if result is None or result.kp_id is None:
            failures.append(
                {
                    "row": row_no,
                    "raw_label": label,
                    "reason": "知识点未匹配到图谱节点，请改用图谱中的标准名称",
                }
            )
            continue

        existing = None
        if batch_id:
            # 明细表内就地编辑后重提该行：按 (batch_id, student_id, kp_id) 覆盖，不新增行
            existing = db.scalar(
                select(ScoreImport).where(
                    ScoreImport.batch_id == batch,
                    ScoreImport.student_id == user.id,
                    ScoreImport.kp_id == result.kp_id,
                )
            )
        if existing is not None:
            existing.score = score
            existing.total = total
            existing.raw_label = label
            existing.source_file = source_file
        else:
            db.add(
                ScoreImport(
                    batch_id=batch,
                    student_id=user.id,
                    kp_id=result.kp_id,
                    score=score,
                    total=total,
                    raw_label=label,
                    source_file=source_file,
                )
            )
        _write_score_alias(db, label, result)
        success.append(
            {
                "row": row_no,
                "raw_label": label,
                "kp_id": result.kp_id,
                "matched_text": result.matched_text,
                "method": result.method,
                "score": score,
                "total": total,
                "updated": existing is not None,
            }
        )

    db.flush()
    profile = learn_profile.refresh_profile_cache(db, user.id)

    return {
        "batch_id": batch,
        "source_file": source_file,
        "success_count": len(success),
        "failure_count": len(failures),
        "success": success,
        "failures": failures,
        "kp_count": len(profile),
    }


def _write_score_alias(db: Session, label: str, result) -> None:
    """成绩导入侧的别名写入（设计文档 3.2.5：三个调用方各自负责写 `kp_alias`）。"""
    norm = normalize_label(label)
    if not norm or result.kp_id is None or result.method == "exact":
        return
    alias = db.scalar(select(KpAlias).where(KpAlias.norm_label == norm))
    if alias is None:
        db.add(
            KpAlias(
                raw_label=label,
                norm_label=norm,
                kp_id=result.kp_id,
                confidence=result.confidence,
                hit_count=1,
                source=SOURCE_SCORE_IMPORT,
            )
        )
    elif alias.source != SOURCE_MANUAL:
        alias.kp_id = result.kp_id
        alias.confidence = result.confidence
        alias.source = SOURCE_SCORE_IMPORT


@router.get("/profile/import-template", summary="下载历史成绩导入模板")
def download_import_template(user: User = Depends(require_student)):
    """模板只含 **知识点 / 得分 / 满分** 三列，**不含学号与姓名**——
    导入的是该生自己的成绩；带上学号反而引出"上传的学号与当前登录用户不一致怎么办"
    这一未定义问题（设计文档 2.2 场景三第 1 条）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "历史成绩"
    sheet.append(list(TEMPLATE_HEADERS))
    sheet.append(["反向传播", 6, 10])
    sheet.append(["梯度下降", 8, 10])
    sheet.append(["卷积神经网络", 7, 10])
    for column, width in zip("ABC", (28, 10, 10)):
        sheet.column_dimensions[column].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    filename = "历史成绩导入模板.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


@router.post("/profile/import", summary="导入历史成绩（文件或 JSON）")
async def import_scores(
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_student),
):
    """同一 handler 接受两种载荷，两条路径复用同一条校验 + 匹配管线。

    JSON 载荷同时承担「手工新增一条」与「明细表内就地编辑后重提该行」两种前端行为。
    """
    content_type = request.headers.get("content-type", "")
    source_file: str | None = None
    batch_id: str | None = None

    if "multipart/form-data" in content_type:
        form = await request.form()
        upload = form.get("file")
        if upload is None or not hasattr(upload, "read"):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="未收到上传文件")
        source_file = getattr(upload, "filename", None)
        payload = await upload.read()
        rows = _parse_score_workbook(payload)
    else:
        try:
            body = ScoreImportRequest.model_validate(await request.json())
        except Exception as exc:  # noqa: BLE001 - 载荷结构错误要给可读提示
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, detail=f"载荷格式不正确：{exc}"
            ) from exc
        batch_id = body.batch_id
        rows = [(row.kp, row.score, row.total) for row in body.rows]

    if not rows:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="没有可导入的数据行")

    report = _import_rows(
        db, user=user, rows=rows, source_file=source_file, batch_id=batch_id
    )
    return ApiResponse.ok(report)


def _parse_score_workbook(payload: bytes) -> list[tuple[str, float, float]]:
    """解析 xlsx：第一行表头，其后每行「知识点 / 得分 / 满分」。"""
    try:
        workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - 非 xlsx 或文件损坏
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail=f"无法解析该文件（需要 xlsx）：{exc}"
        ) from exc

    sheet = workbook.active
    rows: list[tuple[str, float, float]] = []
    for index, raw in enumerate(sheet.iter_rows(values_only=True)):
        if index == 0:
            continue  # 表头
        if raw is None or all(cell is None for cell in raw):
            continue
        cells = list(raw) + [None, None, None]
        label = str(cells[0]).strip() if cells[0] is not None else ""
        try:
            score = float(cells[1]) if cells[1] is not None else None
            total = float(cells[2]) if cells[2] is not None else None
        except (TypeError, ValueError):
            rows.append((label, -1.0, 0.0))  # 交给统一校验产出可读的失败原因
            continue
        rows.append((label, score, total))
    workbook.close()
    return rows


@router.get("/profile/import-batches", summary="导入批次列表")
def list_import_batches(
    db: Session = Depends(get_db),
    user: User = Depends(require_student),
):
    rows = db.execute(
        select(
            ScoreImport.batch_id,
            func.count().label("count"),
            func.min(ScoreImport.created_at).label("created_at"),
            func.max(ScoreImport.source_file).label("source_file"),
        )
        .where(ScoreImport.student_id == user.id)
        .group_by(ScoreImport.batch_id)
        .order_by(func.min(ScoreImport.created_at).desc())
    ).all()

    batch_ids = [row.batch_id for row in rows]
    detail: dict[str, list[dict]] = {batch_id: [] for batch_id in batch_ids}
    if batch_ids:
        names = _kp_names(db)
        for item in db.execute(
            select(ScoreImport)
            .where(ScoreImport.student_id == user.id, ScoreImport.batch_id.in_(batch_ids))
            .order_by(ScoreImport.id)
        ).scalars():
            detail[item.batch_id].append(
                {
                    "id": item.id,
                    "kp_id": item.kp_id,
                    "kp_name": names.get(item.kp_id),
                    "raw_label": item.raw_label,
                    "score": item.score,
                    "total": item.total,
                }
            )

    return ApiResponse.ok(
        {
            "total": len(rows),
            "items": [
                {
                    "batch_id": row.batch_id,
                    "count": row.count,
                    "created_at": row.created_at.isoformat() if row.created_at else None,
                    "source_file": row.source_file,
                    "rows": detail.get(row.batch_id, []),
                }
                for row in rows
            ],
        }
    )


@router.delete("/profile/import-batches/{batch_id}", summary="按批次整批撤销导入")
def delete_import_batch(
    batch_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(require_student),
):
    """撤销后立即重算画像——"撤销了但雷达图没变"是最容易被当场发现的不一致。"""
    result = db.execute(
        delete(ScoreImport).where(
            ScoreImport.batch_id == batch_id, ScoreImport.student_id == user.id
        )
    )
    deleted = result.rowcount or 0
    if not deleted:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="该批次不存在")

    profile = learn_profile.refresh_profile_cache(db, user.id)
    return ApiResponse.ok({"batch_id": batch_id, "deleted": deleted, "kp_count": len(profile)})


# ------------------------------------------------------------------ 教师侧：图谱管理

@router.post("/kp/sync", summary="题库汇入 / 消息挂靠 / 常用问题聚合（可重跑）")
def sync_knowledge(
    target: str = Query("all", pattern="^(questions|messages|faq|all)$"),
    db: Session = Depends(get_db),
    user: User = Depends(require_teacher),
):
    """`questions` 是**演示前置步骤**：不执行它，题库为空，练习页抽不到题。

    **须教师身份**：汇入会重写全库题库的挂靠关系与别名表，是管理动作而非个人数据操作。
    """
    report: dict = {"target": target}
    if target in ("questions", "all"):
        report["questions"] = learn_sync.sync_questions(db)
    if target in ("messages", "all"):
        report["messages"] = learn_sync.sync_messages(db)
    if target in ("faq", "all"):
        report["faq"] = learn_sync.sync_faq(db)
    return ApiResponse.ok(report)


@router.get("/kp/unclassified", summary="未归类知识点清单")
def list_unclassified(
    db: Session = Depends(get_db),
    user: User = Depends(require_teacher),
):
    return ApiResponse.ok(learn_sync.list_unclassified(db))


@router.get("/kp/aliases", summary="全部别名（含误挂改判入口）")
def list_aliases(
    only_placeholder: bool = Query(False),
    db: Session = Depends(get_db),
    user: User = Depends(require_teacher),
):
    """误挂的纠正入口。词典最长子串必然会把"随机梯度下降"挂到"梯度下降"，
    这类错误**不在未归类清单里**（那条记录挂得好好的），只能从全量列表改判。"""
    items = learn_sync.list_aliases(db, only_placeholder=only_placeholder)
    return ApiResponse.ok({"total": len(items), "items": items})


@router.post("/kp/merge", summary="知识点归并 / 误挂改判")
def merge_knowledge(
    body: KpMergeRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_teacher),
):
    try:
        report = learn_sync.merge_alias(
            db, raw_label=body.raw_label, target_kp_id=body.target_kp_id
        )
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    return ApiResponse.ok(report)


__all__ = ["router"]
