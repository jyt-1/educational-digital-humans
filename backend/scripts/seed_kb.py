# [工单18] 人工智能NLP-Agent数字人项目-教育智能体-智能助教任务 —— 公共知识库种子数据脚本
"""一键构建「人工智能专业课程」知识库样本，供工单18 验收演示使用。

对应验收标准第 2 条：「基于人工智能专业课程数据构建公共知识库和个人知识库，
并能够实现资源检索及引用」。脚本会生成 5 份含多模态内容的课程资料
（Word 讲义带表格与图、PDF 讲义含表格页、PPT 课件、Excel 成绩表、示意图），
经**真实的 /api/kb/upload 接口**上传并等待解析完成——走的就是用户点按钮的同一条链路。

用法：
    # 先起后端：cd backend && uvicorn app.main:app --reload
    python scripts/seed_kb.py                      # 用默认账号 demo_teacher/demo123456
    python scripts/seed_kb.py --save-dir ./samples # 顺便把样本文件存下来，便于手工拖拽演示
    python scripts/seed_kb.py --dry-run            # 只生成文件不上传
"""

from __future__ import annotations

import argparse
import io
import sys
import time
from pathlib import Path

import fitz
import httpx
from docx import Document as DocxDocument
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as PptxInches

DOCX_NAME = "人工智能导论-第三章-梯度下降与反向传播.docx"
PDF_NAME = "人工智能导论-第二章-机器学习基础.pdf"
PPTX_NAME = "机器学习实验指导-第一讲.pptx"
XLSX_NAME = "期末知识点分值分布.xlsx"
IMAGE_NAME = "神经网络结构示意图.png"


# ------------------------------------------------------------------ 样本构造

def make_png(color: tuple[int, int, int], size: int = 200) -> bytes:
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, size, size))
    pixmap.set_rect(pixmap.irect, color)
    return pixmap.tobytes("png")


def build_docx() -> bytes:
    document = DocxDocument()
    document.add_heading("第三章 梯度下降与反向传播", level=1)
    document.add_paragraph(
        "梯度下降是一种迭代优化算法：沿着损失函数的负梯度方向更新参数，逐步逼近极小值。"
        "它是神经网络训练的基础，几乎所有深度学习框架都以其为核心优化手段。"
    )
    document.add_heading("3.1 学习率的影响", level=2)
    table = document.add_table(rows=4, cols=3)
    for row, values in enumerate(
        [
            ["学习率", "现象", "建议"],
            ["0.001", "收敛过慢，训练轮次不足时欠拟合", "配合早停策略"],
            ["0.01", "较为平稳，多数任务的首选起点", "常用推荐值"],
            ["0.9", "震荡甚至发散，损失上下跳动", "需衰减或换自适应优化器"],
        ]
    ):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    document.add_paragraph(
        "实践中常令学习率随轮次衰减：η_t = η_0 / (1 + decay × t)，"
        "在训练后期获得更精细的参数更新。"
    )
    document.add_paragraph("下图是不同学习率下损失函数下降曲线的对比示意。")
    document.add_picture(io.BytesIO(make_png((40, 90, 200))), width=Inches(3))
    document.add_heading("3.2 反向传播", level=2)
    document.add_paragraph(
        "反向传播利用链式法则，从输出层向输入层逐层计算梯度：先前向计算出预测值与损失，"
        "再反向把误差梯度分摊到每一层的权重上，最后用梯度下降更新权重。"
    )
    document.add_paragraph("批量梯度下降每次迭代使用全部样本；随机梯度下降每次只用一个样本；"
                           "小批量梯度下降（mini-batch）是两者的折中，也是工程实践的主流。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf() -> bytes:
    doc = fitz.open()
    pages = [
        ("第二章 机器学习基础", "机器学习是让计算机从数据中归纳规律的方法。监督学习使用带标签的数据训练模型。"),
        ("2.1 常见损失函数", "回归任务常用均方误差，分类任务常用交叉熵。损失函数衡量预测与真实的差距。"),
        ("2.2 优化器对照表", "不同优化器在不同任务上的表现差异明显，下表为课程实验中的对照结论。"),
    ]
    for index, (title, body) in enumerate(pages, start=1):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 70), title, fontname="china-s", fontsize=16)
        page.insert_text((72, 110), body, fontname="china-s", fontsize=11)
        if index == 3:
            x0, y0, rows, cols, w, h = 72, 160, 4, 3, 140, 28
            for r in range(rows + 1):
                page.draw_line(fitz.Point(x0, y0 + r * h), fitz.Point(x0 + cols * w, y0 + r * h))
            for c in range(cols + 1):
                page.draw_line(fitz.Point(x0 + c * w, y0), fitz.Point(x0 + c * w, y0 + rows * h))
            cells = [
                ["优化器", "收敛速度", "适用场景"],
                ["SGD", "较慢", "凸问题、小数据"],
                ["Momentum", "中等", "易震荡的损失面"],
                ["Adam", "较快", "深度学习首选"],
            ]
            for r, row in enumerate(cells):
                for c, cell in enumerate(row):
                    page.insert_text(
                        (x0 + c * w + 8, y0 + r * h + 18), cell, fontname="china-s", fontsize=10
                    )
            page.insert_image(fitz.Rect(360, 160, 500, 280), stream=make_png((30, 140, 90)))
    data = doc.tobytes()
    doc.close()
    return data


def build_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "机器学习实验指导：手写数字识别"
    box = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(1.8), PptxInches(7), PptxInches(1.4))
    box.text_frame.text = (
        "实验目标：用三层全连接网络完成 MNIST 手写数字识别；"
        "掌握数据划分、训练循环、准确率评估的完整流程。"
    )
    table_shape = slide.shapes.add_table(4, 3, PptxInches(0.8), PptxInches(3.4), PptxInches(7), PptxInches(2))
    table = table_shape.table
    rows = [
        ["环节", "耗时", "关键点"],
        ["数据加载", "5 分钟", "归一化到 0-1"],
        ["模型训练", "20 分钟", "批量大小 64，轮次 10"],
        ["评估调参", "15 分钟", "混淆矩阵定位错分数字"],
    ]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    slide.shapes.add_picture(
        io.BytesIO(make_png((160, 60, 170), size=160)),
        PptxInches(7.4),
        PptxInches(1.8),
        PptxInches(2.2),
        PptxInches(2.2),
    )
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def build_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "知识点分值分布"
    sheet.append(["知识点", "章节", "分值", "难度"])
    for row in [
        ["梯度下降", "第三章", 15, "中等"],
        ["反向传播", "第三章", 20, "较难"],
        ["过拟合与正则化", "第四章", 15, "中等"],
        ["卷积神经网络", "第五章", 25, "较难"],
        ["模型评估指标", "第二章", 25, "简单"],
    ]:
        sheet.append(row)
    second = workbook.create_sheet("实验成绩")
    second.append(["学号", "姓名", "实验一", "实验二"])
    for row in [["2023001", "李明", 92, 88], ["2023002", "王芳", 85, 91]]:
        second.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_samples(save_dir: Path | None = None) -> dict[str, bytes]:
    samples = {
        DOCX_NAME: build_docx(),
        PDF_NAME: build_pdf(),
        PPTX_NAME: build_pptx(),
        XLSX_NAME: build_xlsx(),
        IMAGE_NAME: make_png((200, 120, 30)),
    }
    if save_dir:
        save_dir.mkdir(parents=True, exist_ok=True)
        for name, data in samples.items():
            (save_dir / name).write_bytes(data)
        print(f"样本文件已保存到 {save_dir}")
    return samples


# ------------------------------------------------------------------ 上传流程

def login_or_register(client: httpx.Client, username: str, password: str) -> str:
    payload = {"username": username, "password": password}
    resp = client.post("/api/auth/login", json=payload)
    if resp.status_code != 200:
        client.post(
            "/api/auth/register",
            json={**payload, "role": "teacher", "display_name": "演示教师"},
        )
        resp = client.post("/api/auth/login", json=payload)
    resp.raise_for_status()
    return resp.json()["data"]["access_token"]


def upload_and_wait(
    client: httpx.Client, headers: dict, name: str, data: bytes, scope: str, timeout: int = 180
) -> dict:
    resp = client.post(
        "/api/kb/upload",
        files={"file": (name, data, "application/octet-stream")},
        data={"scope": scope},
        headers=headers,
    )
    resp.raise_for_status()
    doc = resp.json()["data"]

    deadline = time.time() + timeout
    while time.time() < deadline:
        detail = client.get(f"/api/kb/docs/{doc['id']}", headers=headers).json()["data"]
        if detail["parse_status"] in ("done", "failed"):
            return detail
        time.sleep(1.5)
    raise TimeoutError(f"{name} 解析超时")


def main() -> int:
    parser = argparse.ArgumentParser(description="构建工单18 知识库演示数据")
    parser.add_argument("--base-url", default="http://localhost:8000")
    parser.add_argument("--username", default="demo_teacher")
    parser.add_argument("--password", default="demo123456")
    parser.add_argument("--save-dir", default=None, help="同时把样本文件保存到该目录")
    parser.add_argument("--dry-run", action="store_true", help="只生成文件，不上传")
    parser.add_argument("--private", action="store_true", help="额外往私有库放一份资料")
    args = parser.parse_args()

    save_dir = Path(args.save_dir) if args.save_dir else None
    samples = build_samples(save_dir)
    print(f"已生成 {len(samples)} 份课程资料样本")

    if args.dry_run:
        return 0

    with httpx.Client(base_url=args.base_url, timeout=120.0) as client:
        try:
            token = login_or_register(client, args.username, args.password)
        except Exception as exc:  # noqa: BLE001
            print(f"登录/注册失败：{exc}\n请先启动后端：cd backend && uvicorn app.main:app --reload")
            return 1
        headers = {"Authorization": f"Bearer {token}"}

        print("\n=== 公共知识库（教师上传，全体可检索） ===")
        for name, data in samples.items():
            detail = upload_and_wait(client, headers, name, data, scope="public")
            stats = detail.get("chunk_stats") or {}
            print(
                f"  [{detail['parse_status']}] {name} → {detail['chunk_count']} 块"
                f"（正文 {stats.get('text', 0)} / 表格 {stats.get('table', 0)}"
                f" / 图片 {stats.get('image', 0)} / 公式 {stats.get('formula', 0)}）"
            )
            if detail["parse_status"] == "failed":
                print(f"        失败原因：{detail['parse_error']}")

        if args.private:
            print("\n=== 个人私有知识库（仅本人可见） ===")
            detail = upload_and_wait(
                client, headers, "我的复习笔记-梯度下降.docx", build_docx(), scope="private"
            )
            print(f"  [{detail['parse_status']}] {detail['filename']} → {detail['chunk_count']} 块")

    print("\n完成。可访问前端 /assistant 页面提问，例如：")
    print("  - 梯度下降的学习率过大会有什么后果？")
    print("  - 表格里 Adam 优化器适合什么场景？  （验证表格引用）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
