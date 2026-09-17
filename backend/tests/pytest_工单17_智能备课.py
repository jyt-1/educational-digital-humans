# [工单17] 人工智能NLP-Agent数字人项目-教育智能体-智能备课任务 —— 测试用例
"""工单17 智能备课 pytest 用例。

覆盖：认证与权限、SSE 流式生成、题目入库、列表/详情、编辑保存与版本、
版本回滚、docx/pptx 导出、资源检索、越权隔离、LLM 未配置与 JSON 容错。

LLM 全部 mock，测试不依赖真实 API Key 与网络。
"""

from __future__ import annotations

import io
import json

import pytest
from docx import Document
from pptx import Presentation

from app.services import llm_client
from app.services.llm_client import extract_json_block
from app.services.prompts import build_content_payload, normalize_items

# ============================================================ Mock LLM

LESSON_MD = """# 梯度下降法教案

## 一、教学目标
- **知识目标**：理解梯度下降的数学原理
- **能力目标**：能用 Python 实现批量梯度下降

## 二、教学重点与难点
重点是迭代公式，难点是学习率的选取。

## 三、教学过程
| 教学环节 | 教师活动 | 学生活动 | 时间分配 |
| --- | --- | --- | --- |
| 导入 | 提问最优解 | 思考回答 | 10 分钟 |
| 新知讲授 | 推导公式 | 跟随推导 | 40 分钟 |
"""

SAMPLE_QUESTIONS = [
    {
        "qtype": "单选",
        "stem": "梯度下降中学习率过大最可能导致什么？",
        "options": ["A. 收敛变慢", "B. 震荡不收敛", "C. 梯度消失", "D. 无影响"],
        "answer": "B",
        "analysis": "学习率过大会跨过极小值点，导致震荡甚至发散。",
        "knowledge_point": "梯度下降",
        "difficulty": "简单",
    },
    {
        "qtype": "判断",
        "stem": "批量梯度下降每次迭代使用全部训练样本计算梯度。",
        "options": ["A. 正确", "B. 错误"],
        "answer": "A",
        "analysis": "批量梯度下降的定义即使用全量样本。",
        "knowledge_point": "梯度下降",
        "difficulty": "简单",
    },
]

SAMPLE_SLIDES = [
    {"title": "梯度下降法", "bullets": ["课程：人工智能导论"], "notes": "封面，介绍本节内容"},
    {"title": "核心思想", "bullets": ["沿负梯度方向更新", "逐步逼近极小值"], "notes": "强调直观理解"},
]


def _make_stream(text: str):
    """构造一个假的流式生成器，按 8 字符切片模拟逐字返回。"""

    async def _stream(messages, **kwargs):
        for i in range(0, len(text), 8):
            yield text[i : i + 8]

    return _stream


@pytest.fixture
def mock_llm_markdown(monkeypatch):
    monkeypatch.setattr(llm_client, "chat_stream", _make_stream(LESSON_MD))


@pytest.fixture
def mock_llm_questions(monkeypatch):
    monkeypatch.setattr(
        llm_client, "chat_stream", _make_stream(json.dumps(SAMPLE_QUESTIONS, ensure_ascii=False))
    )


@pytest.fixture
def mock_llm_slides(monkeypatch):
    monkeypatch.setattr(
        llm_client, "chat_stream", _make_stream(json.dumps(SAMPLE_SLIDES, ensure_ascii=False))
    )


def _sse_events(text: str) -> list[tuple[str, dict]]:
    """解析 SSE 响应文本为 [(event, data), ...]。"""
    events = []
    for block in text.strip().split("\n\n"):
        if not block.strip():
            continue
        event, data = None, None
        for line in block.splitlines():
            if line.startswith("event: "):
                event = line[len("event: ") :]
            elif line.startswith("data: "):
                data = json.loads(line[len("data: ") :])
        if event:
            events.append((event, data))
    return events


def _generate(client, auth, token, content_type: str, **overrides):
    """调用生成接口并返回 (events, done_data)。"""
    payload = {
        "content_type": content_type,
        "course_name": "人工智能导论",
        "subject": "人工智能",
        "chapter": "第3章 神经网络",
        "knowledge_points": ["梯度下降", "反向传播"],
        "difficulty": "中等",
        "objectives": ["理解梯度下降"],
    }
    payload.update(overrides)
    resp = client.post("/api/lesson/generate", json=payload, headers=auth(token))
    assert resp.status_code == 200, resp.text
    events = _sse_events(resp.text)
    done = next((d for e, d in events if e == "done"), None)
    return events, done


# ============================================================ 1. 认证与权限

class TestAuth:
    def test_health(self, client):
        resp = client.get("/api/health")
        assert resp.status_code == 200
        body = resp.json()
        assert body["code"] == 0
        assert "llm_configured" in body["data"]

    def test_register_returns_user_without_password(self, client):
        resp = client.post(
            "/api/auth/register",
            json={"username": "t_new", "password": "pwd123456", "role": "teacher"},
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()["data"]
        assert data["username"] == "t_new"
        assert data["role"] == "teacher"
        # 响应中绝不能出现密码或哈希
        assert "password" not in json.dumps(data)

    def test_register_duplicate_username(self, client):
        client.post(
            "/api/auth/register",
            json={"username": "dup_user", "password": "pwd123456", "role": "student"},
        )
        resp = client.post(
            "/api/auth/register",
            json={"username": "dup_user", "password": "pwd123456", "role": "student"},
        )
        assert resp.status_code == 409

    def test_register_rejects_invalid_role(self, client):
        resp = client.post(
            "/api/auth/register",
            json={"username": "bad_role", "password": "pwd123456", "role": "hacker"},
        )
        assert resp.status_code == 400

    def test_login_wrong_password(self, client, teacher_token):
        resp = client.post(
            "/api/auth/login", json={"username": "teacher_zhang", "password": "wrong-pwd"}
        )
        assert resp.status_code == 401
        # 与"用户不存在"返回同一提示，避免账号枚举
        resp2 = client.post(
            "/api/auth/login", json={"username": "no_such_user", "password": "whatever"}
        )
        assert resp2.status_code == 401
        assert resp.json()["msg"] == resp2.json()["msg"]

    def test_me_returns_current_user(self, client, auth, teacher_token):
        resp = client.get("/api/auth/me", headers=auth(teacher_token))
        assert resp.status_code == 200
        assert resp.json()["data"]["username"] == "teacher_zhang"

    def test_password_stored_as_hash(self, client):
        """密码必须以 bcrypt 哈希存储，不得明文。"""
        from sqlalchemy import select

        from app.db import SessionLocal
        from app.models.user import User

        with SessionLocal() as db:
            user = db.scalar(select(User).where(User.username == "teacher_zhang"))
            assert user is not None
            assert user.password_hash != "pwd123456"
            assert user.password_hash.startswith("$2")

    def test_generate_requires_auth(self, client):
        resp = client.post(
            "/api/lesson/generate",
            json={"content_type": "教案", "course_name": "人工智能导论"},
        )
        assert resp.status_code == 401

    def test_student_cannot_access_lesson(self, client, auth, student_token):
        resp = client.get("/api/lesson/plans", headers=auth(student_token))
        assert resp.status_code == 403

    def test_invalid_token_rejected(self, client):
        resp = client.get("/api/lesson/plans", headers={"Authorization": "Bearer not-a-jwt"})
        assert resp.status_code == 401


# ============================================================ 2. SSE 流式生成

class TestGenerate:
    def test_generate_lesson_plan_streams_and_saves(self, client, auth, teacher_token, mock_llm_markdown):
        events, done = _generate(client, auth, teacher_token, "教案")

        # 有增量事件，能拼回完整内容
        deltas = [d["text"] for e, d in events if e == "delta"]
        assert len(deltas) > 1, "应当分多次推送增量"
        assert "".join(deltas) == LESSON_MD

        assert done is not None and done["plan_id"] > 0
        assert done["version"] == 1

        plan_id = done["plan_id"]
        resp = client.get(f"/api/lesson/plans/{plan_id}", headers=auth(teacher_token))
        assert resp.status_code == 200
        detail = resp.json()["data"]
        assert detail["content_type"] == "教案"
        assert detail["content"]["format"] == "markdown"
        assert "梯度下降" in detail["content"]["raw"]
        assert detail["knowledge_points"] == ["梯度下降", "反向传播"]
        assert detail["current_version"] == 1

    def test_generate_exercises_persists_questions(self, client, auth, teacher_token, mock_llm_questions):
        events, done = _generate(client, auth, teacher_token, "习题")
        assert done["question_count"] == len(SAMPLE_QUESTIONS)

        from sqlalchemy import select

        from app.db import SessionLocal
        from app.models.lesson import Exercise

        with SessionLocal() as db:
            rows = db.scalars(
                select(Exercise).where(Exercise.plan_id == done["plan_id"])
            ).all()
            assert len(rows) == len(SAMPLE_QUESTIONS)
            assert rows[0].stem == SAMPLE_QUESTIONS[0]["stem"]
            assert rows[0].knowledge_point == "梯度下降"
            assert json.loads(rows[0].options_json) == SAMPLE_QUESTIONS[0]["options"]

    def test_generate_exam_persists_questions_with_score(self, client, auth, teacher_token, mock_llm_questions):
        events, done = _generate(client, auth, teacher_token, "试题")

        from sqlalchemy import select

        from app.db import SessionLocal
        from app.models.lesson import ExamQuestion

        with SessionLocal() as db:
            rows = db.scalars(
                select(ExamQuestion).where(ExamQuestion.plan_id == done["plan_id"])
            ).all()
            assert len(rows) == len(SAMPLE_QUESTIONS)

    def test_generate_courseware_keeps_slides(self, client, auth, teacher_token, mock_llm_slides):
        events, done = _generate(client, auth, teacher_token, "课件")
        resp = client.get(f"/api/lesson/plans/{done['plan_id']}", headers=auth(teacher_token))
        content = resp.json()["data"]["content"]
        assert content["format"] == "json"
        assert len(content["items"]) == len(SAMPLE_SLIDES)
        assert content["items"][0]["title"] == "梯度下降法"

    def test_generate_rejects_invalid_content_type(self, client, auth, teacher_token):
        resp = client.post(
            "/api/lesson/generate",
            json={"content_type": "小说", "course_name": "人工智能导论"},
            headers=auth(teacher_token),
        )
        assert resp.status_code == 400

    def test_generate_reports_error_when_llm_not_configured(self, client, auth, teacher_token):
        """未配置 API Key 时应以 SSE error 事件告知，而不是 500。"""
        events, done = _generate(client, auth, teacher_token, "教案")
        error = next((d for e, d in events if e == "error"), None)
        assert done is None
        assert error is not None
        assert "API Key" in error["msg"]

    def test_generate_survives_llm_exception(self, client, auth, teacher_token, monkeypatch):
        async def _boom(messages, **kwargs):
            yield "部分内容"
            raise RuntimeError("上游超时")

        monkeypatch.setattr(llm_client, "chat_stream", _boom)
        events, done = _generate(client, auth, teacher_token, "教案")
        error = next((d for e, d in events if e == "error"), None)
        assert error is not None and "上游超时" in error["msg"]
        assert done is None


# ============================================================ 3. 列表与详情

class TestListDetail:
    def test_list_returns_own_plans(self, client, auth, teacher_token, mock_llm_markdown):
        _generate(client, auth, teacher_token, "教案")
        resp = client.get("/api/lesson/plans", headers=auth(teacher_token))
        assert resp.status_code == 200
        assert len(resp.json()["data"]) >= 1

    def test_list_filter_by_type(self, client, auth, teacher_token, mock_llm_markdown):
        _generate(client, auth, teacher_token, "教案")
        resp = client.get(
            "/api/lesson/plans", params={"content_type": "课件"}, headers=auth(teacher_token)
        )
        assert resp.status_code == 200
        assert all(r["content_type"] == "课件" for r in resp.json()["data"])

    def test_stats_groups_by_type(self, client, auth, teacher_token, mock_llm_markdown):
        _generate(client, auth, teacher_token, "教案")
        resp = client.get("/api/lesson/stats", headers=auth(teacher_token))
        assert resp.status_code == 200
        assert resp.json()["data"]["by_type"]["教案"] >= 1

    def test_detail_404_for_other_teacher(self, client, auth, teacher_token, other_teacher_token, mock_llm_markdown):
        """越权隔离：他人备课记录一律 404。"""
        _, done = _generate(client, auth, teacher_token, "教案")
        resp = client.get(
            f"/api/lesson/plans/{done['plan_id']}", headers=auth(other_teacher_token)
        )
        assert resp.status_code == 404

    def test_detail_404_for_missing_plan(self, client, auth, teacher_token):
        resp = client.get("/api/lesson/plans/999999", headers=auth(teacher_token))
        assert resp.status_code == 404


# ============================================================ 4. 编辑与版本管理

class TestVersioning:
    def test_update_creates_new_version(self, client, auth, teacher_token, mock_llm_markdown):
        _, done = _generate(client, auth, teacher_token, "教案")
        plan_id = done["plan_id"]

        edited = {"format": "markdown", "raw": "# 教师修改后的教案\n\n新增内容。"}
        resp = client.put(
            f"/api/lesson/plans/{plan_id}",
            json={"content": edited, "remark": "教师补充案例"},
            headers=auth(teacher_token),
        )
        assert resp.status_code == 200, resp.text
        detail = resp.json()["data"]
        assert detail["current_version"] == 2
        assert detail["content"]["raw"] == "# 教师修改后的教案\n\n新增内容。"

        versions = client.get(f"/api/lesson/plans/{plan_id}/versions", headers=auth(teacher_token))
        assert versions.status_code == 200
        vlist = versions.json()["data"]
        assert len(vlist) == 2
        assert vlist[0]["version_no"] == 2
        assert vlist[0]["remark"] == "教师补充案例"
        assert vlist[1]["version_no"] == 1

    def test_rollback_restores_old_content_without_losing_history(
        self, client, auth, teacher_token, mock_llm_markdown
    ):
        _, done = _generate(client, auth, teacher_token, "教案")
        plan_id = done["plan_id"]

        client.put(
            f"/api/lesson/plans/{plan_id}",
            json={"content": {"format": "markdown", "raw": "改坏了的内容"}},
            headers=auth(teacher_token),
        )
        versions = client.get(
            f"/api/lesson/plans/{plan_id}/versions", headers=auth(teacher_token)
        ).json()["data"]
        v1_id = next(v["id"] for v in versions if v["version_no"] == 1)

        resp = client.post(
            f"/api/lesson/plans/{plan_id}/versions/{v1_id}/rollback",
            headers=auth(teacher_token),
        )
        assert resp.status_code == 200, resp.text
        detail = resp.json()["data"]
        # 内容回到 v1
        assert "梯度下降" in detail["content"]["raw"]
        # 回滚不删历史，而是追加新版本
        assert detail["current_version"] == 3
        after = client.get(
            f"/api/lesson/plans/{plan_id}/versions", headers=auth(teacher_token)
        ).json()["data"]
        assert len(after) == 3
        assert after[0]["remark"] == "回滚自 v1"

    def test_rollback_404_for_foreign_version(self, client, auth, teacher_token, other_teacher_token, mock_llm_markdown):
        _, done = _generate(client, auth, teacher_token, "教案")
        resp = client.post(
            f"/api/lesson/plans/{done['plan_id']}/versions/999999/rollback",
            headers=auth(other_teacher_token),
        )
        assert resp.status_code == 404

    def test_editing_questions_refreshes_question_table(self, client, auth, teacher_token, mock_llm_questions):
        """编辑习题后，工单19 抽题用的 exercises 表应同步为最新内容。"""
        _, done = _generate(client, auth, teacher_token, "习题")
        plan_id = done["plan_id"]

        new_items = [dict(SAMPLE_QUESTIONS[0], stem="修改后的题干")]
        client.put(
            f"/api/lesson/plans/{plan_id}",
            json={"content": {"format": "json", "raw": "x", "items": new_items}},
            headers=auth(teacher_token),
        )

        from sqlalchemy import select

        from app.db import SessionLocal
        from app.models.lesson import Exercise

        with SessionLocal() as db:
            rows = db.scalars(select(Exercise).where(Exercise.plan_id == plan_id)).all()
            assert len(rows) == 1
            assert rows[0].stem == "修改后的题干"


# ============================================================ 5. 导出

class TestExport:
    def test_export_lesson_plan_docx_matches_edit(self, client, auth, teacher_token, mock_llm_markdown):
        """验收要求：导出内容与编辑后的内容一致。"""
        _, done = _generate(client, auth, teacher_token, "教案")
        plan_id = done["plan_id"]

        marker = "教师特意添加的验收标记文字"
        client.put(
            f"/api/lesson/plans/{plan_id}",
            json={"content": {"format": "markdown", "raw": f"# 教案\n\n{marker}\n"}},
            headers=auth(teacher_token),
        )

        resp = client.get(
            f"/api/lesson/plans/{plan_id}/export",
            params={"format": "docx"},
            headers=auth(teacher_token),
        )
        assert resp.status_code == 200
        assert resp.content[:2] == b"PK"  # docx 是 zip
        doc = Document(io.BytesIO(resp.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert marker in text

    def test_export_extracts_tables_to_docx(self, client, auth, teacher_token, mock_llm_markdown):
        """教案中的 Markdown 表格应渲染为 docx 表格。"""
        _, done = _generate(client, auth, teacher_token, "教案")
        resp = client.get(
            f"/api/lesson/plans/{done['plan_id']}/export",
            params={"format": "docx"},
            headers=auth(teacher_token),
        )
        doc = Document(io.BytesIO(resp.content))
        assert len(doc.tables) >= 1
        header = [c.text for c in doc.tables[0].rows[0].cells]
        assert "教学环节" in header

    def test_export_exercises_docx_contains_answers(self, client, auth, teacher_token, mock_llm_questions):
        _, done = _generate(client, auth, teacher_token, "习题")
        resp = client.get(
            f"/api/lesson/plans/{done['plan_id']}/export",
            params={"format": "docx"},
            headers=auth(teacher_token),
        )
        assert resp.status_code == 200
        doc = Document(io.BytesIO(resp.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "【答案】" in text
        assert "【解析】" in text
        assert SAMPLE_QUESTIONS[0]["stem"] in text

    def test_export_courseware_pptx(self, client, auth, teacher_token, mock_llm_slides):
        _, done = _generate(client, auth, teacher_token, "课件")
        resp = client.get(
            f"/api/lesson/plans/{done['plan_id']}/export",
            params={"format": "pptx"},
            headers=auth(teacher_token),
        )
        assert resp.status_code == 200
        assert resp.content[:2] == b"PK"
        prs = Presentation(io.BytesIO(resp.content))
        # 封面 + 每页 slide
        assert len(prs.slides) == len(SAMPLE_SLIDES) + 1
        titles = [s.shapes.title.text for s in prs.slides]
        assert "梯度下降法" in titles

    def test_export_pptx_rejected_for_non_courseware(self, client, auth, teacher_token, mock_llm_markdown):
        _, done = _generate(client, auth, teacher_token, "教案")
        resp = client.get(
            f"/api/lesson/plans/{done['plan_id']}/export",
            params={"format": "pptx"},
            headers=auth(teacher_token),
        )
        assert resp.status_code == 400

    def test_export_filename_is_utf8_encoded(self, client, auth, teacher_token, mock_llm_markdown):
        _, done = _generate(client, auth, teacher_token, "教案")
        resp = client.get(
            f"/api/lesson/plans/{done['plan_id']}/export", headers=auth(teacher_token)
        )
        disposition = resp.headers["content-disposition"]
        assert "filename*=UTF-8''" in disposition

    def test_export_404_for_other_teacher(self, client, auth, teacher_token, other_teacher_token, mock_llm_markdown):
        _, done = _generate(client, auth, teacher_token, "教案")
        resp = client.get(
            f"/api/lesson/plans/{done['plan_id']}/export", headers=auth(other_teacher_token)
        )
        assert resp.status_code == 404


# ============================================================ 6. 资源检索

class TestResources:
    def test_search_resources_by_keyword(self, client, auth, teacher_token):
        from app.db import SessionLocal
        from app.models.lesson import Resource

        with SessionLocal() as db:
            db.add(
                Resource(
                    title="校本讲义-梯度下降",
                    source="校本",
                    content="本讲义介绍梯度下降的直观理解。",
                    is_public=True,
                )
            )
            db.commit()

        resp = client.post(
            "/api/lesson/resources/search",
            json={"keyword": "梯度下降"},
            headers=auth(teacher_token),
        )
        assert resp.status_code == 200
        titles = [r["title"] for r in resp.json()["data"]]
        assert "校本讲义-梯度下降" in titles

    def test_search_returns_empty_for_no_match(self, client, auth, teacher_token):
        resp = client.post(
            "/api/lesson/resources/search",
            json={"keyword": "一个绝对不存在的关键词xyz"},
            headers=auth(teacher_token),
        )
        assert resp.status_code == 200
        assert resp.json()["data"] == []


# ============================================================ 7. JSON 容错（单位测试）

class TestJsonTolerance:
    def test_extract_plain_json(self):
        assert extract_json_block('{"a": 1}') == {"a": 1}

    def test_extract_from_fenced_code_block(self):
        text = '好的，以下是结果：\n```json\n{"a": 1}\n```\n希望有帮助。'
        assert extract_json_block(text) == {"a": 1}

    def test_extract_from_prose_without_fence(self):
        text = '这是题目：{"questions": [{"stem": "题干"}]} 请查收。'
        assert extract_json_block(text) == {"questions": [{"stem": "题干"}]}

    def test_extract_array(self):
        assert extract_json_block('说明文字 [{"a": 1}, {"b": 2}] 结束') == [{"a": 1}, {"b": 2}]

    def test_extract_handles_braces_inside_strings(self):
        text = '{"stem": "代码里的 { 符号 } 不应破坏配对", "answer": "A"}'
        parsed = extract_json_block(text)
        assert parsed["answer"] == "A"

    def test_extract_returns_none_on_garbage(self):
        assert extract_json_block("完全不是 JSON 的一段话") is None
        assert extract_json_block("") is None

    def test_normalize_items_accepts_wrapped_shape(self):
        assert normalize_items({"questions": [{"stem": "a"}]}) == [{"stem": "a"}]
        assert normalize_items([{"stem": "a"}]) == [{"stem": "a"}]
        assert normalize_items({"unknown": 1}) == []

    def test_build_content_payload_marks_format(self):
        md = json.loads(build_content_payload("教案", "# 标题", None))
        assert md["format"] == "markdown" and "items" not in md

        js = json.loads(build_content_payload("习题", "raw", [{"stem": "a"}]))
        assert js["format"] == "json" and js["items"] == [{"stem": "a"}]


# ============================================================ 8. 导出渲染（单位测试）

class TestExporterUnit:
    def test_questions_to_docx_handles_missing_fields(self):
        from app.services.exporter import questions_to_docx

        # 残缺数据不应导致导出崩溃
        data = questions_to_docx([{"stem": "只有题干"}, {}], "题目")
        assert data[:2] == b"PK"

    def test_courseware_to_pptx_handles_empty_bullets(self):
        from app.services.exporter import courseware_to_pptx

        data = courseware_to_pptx([{"title": "空页", "bullets": []}], "课件")
        assert data[:2] == b"PK"
        prs = Presentation(io.BytesIO(data))
        assert len(prs.slides) == 2

    def test_markdown_to_docx_bold_and_lists(self):
        from app.services.exporter import markdown_to_docx

        md = "# 标题\n\n- 项目一\n- **加粗项**\n\n普通段落。"
        data = markdown_to_docx(md, "测试")
        doc = Document(io.BytesIO(data))
        texts = [p.text for p in doc.paragraphs]
        assert any("加粗项" in t for t in texts)
        assert any(r.bold for p in doc.paragraphs for r in p.runs if r.text == "加粗项")


# ============================================================ 9. 提示词输出规格

class TestPromptSchema:
    """守住输出规格：模板没要求的字段，LLM 就不会给。

    回归背景：习题模板原先只要求 qtype/stem/options/answer/analysis/knowledge_point，
    实测 23 道题 score 全为 None，教师得手工补 23 次分值。
    """

    @pytest.mark.parametrize("content_type", ["习题", "试题"])
    def test_question_prompt_requires_score(self, content_type):
        from app.services.prompts import _USER_TEMPLATES

        template = _USER_TEMPLATES[content_type]
        assert '"score"' in template, f"{content_type}模板必须要求 LLM 输出 score 字段"

    @pytest.mark.parametrize("content_type", ["习题", "试题"])
    def test_question_prompt_requires_core_fields(self, content_type):
        from app.services.prompts import _USER_TEMPLATES

        template = _USER_TEMPLATES[content_type]
        for field in ("qtype", "stem", "options", "answer", "analysis", "knowledge_point"):
            assert f'"{field}"' in template, f"{content_type}模板缺少字段 {field}"

    def test_courseware_prompt_requires_bullets_and_notes(self):
        from app.services.prompts import _USER_TEMPLATES

        template = _USER_TEMPLATES["课件"]
        # 幻灯片编辑器与 pptx 导出都依赖这三个字段
        for field in ("title", "bullets", "notes"):
            assert f'"{field}"' in template

    def test_every_content_type_has_template(self):
        from app.services.prompts import _USER_TEMPLATES
        from app.models.lesson import CONTENT_TYPES

        assert set(_USER_TEMPLATES) == set(CONTENT_TYPES)

    def test_exam_spec_sums_to_100(self):
        """月考试题分项相加必须恰为 100 分。

        回归背景：模板原先手写「总分 100 分」，但列的分值是 10×3+5×4+5×2+3×10=90，
        LLM 按分项出题得到 90 分的卷子，与卷面声明的总分对不上。
        """
        from app.services.prompts import DEFAULT_EXAM_SPEC

        assert sum(n * s for n, s in DEFAULT_EXAM_SPEC.values()) == 100

    def test_exercise_spec_matches_work_order(self):
        """工单17 规定习题题量：单选10 + 多选5 + 判断5 + 简答3。"""
        from app.services.prompts import DEFAULT_EXERCISE_SPEC

        assert DEFAULT_EXERCISE_SPEC == {"单选": (10, 2), "多选": (5, 4), "判断": (5, 2), "简答": (3, 10)}
        assert sum(n * s for n, s in DEFAULT_EXERCISE_SPEC.values()) == 80

    def test_spec_line_reports_consistent_totals(self):
        from app.services.prompts import _spec_line, DEFAULT_EXAM_SPEC

        line = _spec_line(DEFAULT_EXAM_SPEC)
        assert "总分 100 分" in line
        assert "共 23 道" in line

    @pytest.mark.parametrize(
        ("content_type", "expect_total"),
        [("习题", "总分 80 分"), ("试题", "总分 100 分")],
    )
    def test_built_prompt_states_correct_total(self, content_type, expect_total):
        """真正送给 LLM 的 prompt 里，声明的总分必须与分项一致。"""
        from app.services.prompts import build_messages

        messages = build_messages(
            content_type,
            subject="人工智能",
            course_name="人工智能导论",
            chapter="第3章",
            knowledge_points=["反向传播"],
            difficulty="中等",
            objectives=[],
        )
        assert expect_total in messages[1]["content"]

    # ---- 设计文档 3.2.6（v1.3）：知识点与难度的取值约束 ----
    #
    # 回归背景：工单19 的知识点匹配采用「后缀归一化」吃掉标签漂移，其前提是
    # LLM 从上下文【涉及知识点】里原样选词，而非自己改写（"BP算法" vs "反向传播"）。
    # 模板一旦松口，未归类节点就会被撑爆，匹配方案随之失效。

    @pytest.mark.parametrize("content_type", ["习题", "试题"])
    def test_question_prompt_pins_knowledge_point_source(self, content_type):
        from app.services.prompts import _USER_TEMPLATES

        template = _USER_TEMPLATES[content_type]
        assert "【涉及知识点】" in template, f"{content_type}模板须指明 knowledge_point 的来源"
        assert "原样选取" in template, f"{content_type}模板须禁止改写知识点名称"

    @pytest.mark.parametrize("content_type", ["习题", "试题"])
    def test_question_prompt_pins_difficulty_enum(self, content_type):
        """difficulty 只能取三档，与工单19 的 CHECK 约束同口径。"""
        from app.services.prompts import _USER_TEMPLATES

        template = _USER_TEMPLATES[content_type]
        assert "「简单」「中等」「困难」" in template, f"{content_type}模板须给出 difficulty 三档枚举"

    def test_exam_prompt_requires_knowledge_point_in_requirements(self):
        """回归：试题模板的「要求」区块原先漏了 knowledge_point。

        示例 JSON 里虽写了该字段，但要求列表没提，LLM 实测会大面积为空——
        而试题正是工单19 题库的主要来源。
        """
        from app.services.prompts import _TEMPLATE_EXAM

        requirements = _TEMPLATE_EXAM.split("要求：", 1)[1]
        assert "knowledge_point" in requirements

    @pytest.mark.parametrize("content_type", ["习题", "试题"])
    def test_template_example_item_matches_downstream_schema(self, content_type):
        """把模板里的示例 JSON 当 fixture 解析，断言它满足工单19 汇入所需的全部字段。

        这样"模板承诺的输出结构"与"下游消费方要的结构"被钉在同一个断言里：
        谁改坏了示例块（例如删掉 knowledge_point），这条用例立刻红。
        """
        from app.services.prompts import _USER_TEMPLATES

        items = _example_items(_USER_TEMPLATES[content_type])
        assert len(items) == 1, f"{content_type}模板示例应为单道题的结构样例"

        item = items[0]
        for field in ("qtype", "stem", "options", "answer", "analysis", "knowledge_point", "score", "difficulty"):
            assert field in item, f"{content_type}模板示例缺少字段 {field}"

        assert item["knowledge_point"], f"{content_type}模板示例的 knowledge_point 不得为空"
        assert item["difficulty"] in {"简单", "中等", "困难"}
        assert isinstance(item["score"], int) and item["score"] > 0

        # 解析路径必须原样接受该结构
        assert normalize_items([item]) == [item]


def _example_items(template: str) -> list[dict]:
    """抽出模板「每道题的结构如下：」之后的示例 JSON 数组。

    模板是 .format() 字符串，花括号被转义成 {{ }}，故先反转义再按括号深度配对，
    避免误取到「要求」区块里 `["A. 正确", "B. 错误"]` 之类的内层数组。
    """
    raw = template.replace("{{", "{").replace("}}", "}")
    start = raw.index("[", raw.index("每道题的结构如下："))

    depth = 0
    for idx in range(start, len(raw)):
        if raw[idx] == "[":
            depth += 1
        elif raw[idx] == "]":
            depth -= 1
            if depth == 0:
                return json.loads(raw[start : idx + 1])

    raise AssertionError("模板中未找到闭合的示例 JSON 数组")
